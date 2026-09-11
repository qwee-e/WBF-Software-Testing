"""Fill only member A rows in a COPY of the supplied workbook from real evidence.

Reporting dependency: openpyxl==3.1.5 (separate from the tested .venv).
Run from the repository: python -X utf8 manual_tests/member_a/build_report.py
"""
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "test_data/member_a"
OUT = ROOT / "docs/member_a"
ASSIGNED = [1, 3, 5, 7, 9, 11, 13, 15, 17, 19, 21, 23, 25, 26, 29]


def compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def main():
    evidence = {}
    for path in sorted((ROOT / "manual_tests/member_a/results").glob("batch_*/results.json")):
        batch = json.loads(path.read_text(encoding="utf-8"))
        for result in batch["results"]:
            if result["case_id"] in evidence:
                raise ValueError("Duplicate initial execution evidence")
            evidence[result["case_id"]] = (result, batch, path.relative_to(ROOT).as_posix())
    workbook = openpyxl.load_workbook(DATA / "source_case_list.xlsx")
    sheet = workbook.worksheets[1]
    summary = workbook.create_sheet("成员一执行汇总")
    summary.append(["编号", "标题", "状态", "执行时间", "源码提交（运行前）", "证据路径"])
    markdown = ["# 成员一 WBF 模块1测试执行记录", "",
                "按分工执行 01、03、05、07、09、11、13、15、17、19、21、23、25、26、29。其他成员的结果不代填。",
                "输入来自用户提供的附件，预期在执行前独立填写，实际结果来自 pytest 记录；这不是官方测试复现材料。", "",
                "## 本次统计", ""]
    counts = {s: sum(r[0]["status"] == s for r in evidence.values()) for s in ("OK", "NG", "NT")}
    markdown += [f"已执行 {len(evidence)}/15：OK {counts['OK']}，NG {counts['NG']}，NT {counts['NT']}；待执行 {15-len(evidence)}。", "",
                 "NG 保留为真实失败，不使用 xfail 隐藏。23、25 的 ValueError 预期属于建议健壮性契约，缺陷定性需评审；26 的有限坐标要求用于检查数值有效性。", "",
                 "## 逐项结果", "", "| 编号 | 测试内容 | 状态 | 执行证据 |", "| --- | --- | --- | --- |"]
    for number in ASSIGNED:
        case_id = f"WBF-UT-{number:03d}"
        row = next(r for r in range(2, 32) if sheet.cell(r, 1).value == case_id)
        title = sheet.cell(row, 3).value
        if case_id not in evidence:
            summary.append([case_id, title, "待执行"])
            markdown.append(f"| {case_id} | {title} | 待执行 | — |")
            continue
        result, batch, path = evidence[case_id]
        spec = json.loads((DATA / f"{case_id}.json").read_text(encoding="utf-8"))
        details = []
        for call in result["calls"]:
            details.append(compact({k: call[k] for k in ("scenario", "output", "shapes", "exception", "warnings", "stdout")}))
        actual = "\n".join(details)
        sheet.cell(row, 9, spec["procedure"])
        sheet.cell(row, 10, spec["expected_description"])
        sheet.cell(row, 11, actual)
        sheet.cell(row, 12, result["status"])
        sheet.cell(row, 13, spec["method"] + "；" + spec["notes"] +
                   f"\n执行时间：{batch['executed_at']}；Python {batch['python']} / NumPy {batch['numpy']} / pytest {batch['pytest']}" +
                   f"\n运行前源码提交：{batch['source_commit']}；WBF 与 reproduction-success 一致：{batch['matches_reproduction_success']}；证据：{path}")
        for col in range(9, 14):
            sheet.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = 120
        sheet.cell(row, 12).fill = PatternFill("solid", fgColor="C6EFCE" if result["status"] == "OK" else "FFC7CE")
        summary.append([case_id, title, result["status"], batch["executed_at"], batch["source_commit"], path])
        markdown.append(f"| {case_id} | {title} | {result['status']} | [JSON](../../{path}) / [日志](../../{path.replace('results.json','pytest.log')}) |")
    info = workbook.worksheets[0]
    info["E7"] = "reproduction-success（WBF源码哈希逐批核对）"
    info["E8"] = "WBF 模块1测试"
    info["E9"] = "lry（成员一；工具辅助编写与执行，待成员复核）"
    info["J9"] = max((b[1]["executed_at"] for b in evidence.values()), default="待执行")
    # Fix template's 65535-COUNTBLANK(A:A) formula; keep total scope at 30 source cases.
    info["E13"] = f"=COUNTA('{sheet.title}'!A2:A31)"
    info["B16"] = "保留原清单30条输入，仅回填成员一15条的步骤、预期、实际结果、状态和备注。空白结果不表示通过。详见成员一执行汇总。"
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = f"A1:F{summary.max_row}"
    for col, width in {"A": 18, "B": 52, "C": 12, "D": 32, "E": 45, "F": 65}.items():
        summary.column_dimensions[col].width = width
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / "成员一_WBF模块1测试用例与结果.xlsx"
    workbook.save(target)
    markdown += ["", "## 详细记录", ""]
    for case_id, (result, batch, path) in evidence.items():
        spec = json.loads((DATA / f"{case_id}.json").read_text(encoding="utf-8"))
        markdown += [f"### {case_id} {spec['title']}", "", f"- 预期：{spec['expected_description']}",
                     f"- 判定：**{result['status']}**；方法：{spec['method']}", f"- 备注：{spec['notes']}",
                     f"- 时间：{batch['executed_at']}；被测文件 SHA256：`{batch['source_sha256']}`。", "", "```json",
                     json.dumps(result["calls"], ensure_ascii=False, indent=2), "```", ""]
        if result["failure"]:
            markdown += ["断言失败原文：", "", "```text", result["failure"], "```", ""]
    markdown += ["## 报告使用说明", "",
                 "Excel 是原附件的结果副本，未改动下载目录中的原文件；A:H 原用例信息保持不变，仅成员一行回填 I:M，并新增汇总页。表格模板中的通用说明属于附件内容，不替代用户本次记录执行结果的要求。",
                 "文档由 build_report.py 从真实 JSON 记录生成；生成工具使用单独环境的 openpyxl 3.1.5，不改变被测虚拟环境。"]
    (OUT / "执行记录.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    # Read the delivered file back: source columns and other member's result cells must be unchanged.
    check = openpyxl.load_workbook(target)
    original = openpyxl.load_workbook(DATA / "source_case_list.xlsx")
    for row in range(2, 32):
        for col in range(1, 9):
            assert check.worksheets[1].cell(row, col).value == original.worksheets[1].cell(row, col).value
        if int(str(sheet.cell(row, 1).value).split("-")[-1]) not in ASSIGNED:
            for col in range(9, 14):
                assert check.worksheets[1].cell(row, col).value == original.worksheets[1].cell(row, col).value
    print(f"Report verified: {len(evidence)}/15 executed; {counts}; source inputs and member B rows preserved.")


if __name__ == "__main__":
    main()
